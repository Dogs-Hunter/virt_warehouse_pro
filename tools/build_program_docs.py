from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Game\Test")
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)

BLACK = "000000"
NAVY = "17365D"
PALE = "EDF3F8"
GRAY = "D9D9D9"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), GRAY)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def base_document(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(9)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 16, 14, 6), ("Heading 2", 12.5, 10, 4), ("Heading 3", 11, 8, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")

    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(70, 70, 70)

    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    return doc


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 24 and i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(9.2)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_architecture(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(
        "Клиент  →  HAProxy  →  два приложения Rust  →  NATS JetStream 3 узла\n"
        "                                    ↓\n"
        "                         WAL  +  Snapshot  +  Redb"
    )
    r.font.name = "Cascadia Mono"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cascadia Mono")
    r.font.size = Pt(9.5)


def build_short():
    doc = base_document(
        "Краткое описание виртуального склада",
        "Назначение, архитектура, обработка запросов и основные гарантии",
    )
    p = doc.add_paragraph()
    p.add_run("Назначение. ").bold = True
    p.add_run(
        "Программа ведёт складские остатки по паре владелец и товар, принимает изменения "
        "через HTTP API и не допускает повторного применения одной бизнес операции. "
        "Решение рассчитано на быстрые пакетные запросы и продолжение работы после отказа одного узла."
    )

    add_heading(doc, "Состав системы", 1)
    add_architecture(doc)
    add_table(doc, ["Компонент", "Назначение"], [
        ("HAProxy 3.2", "Принимает HTTPS трафик и переключает запросы между приложениями"),
        ("Rust Axum Tokio", "Проверяет, записывает и читает складские операции"),
        ("NATS JetStream", "Хранит канонический журнал на трёх репликах с кворумом 2 из 3"),
        ("WAL и Snapshot", "Ускоряют локальное восстановление приложения"),
        ("Redb", "Хранит дедупликацию, историю и индексы чтения"),
        ("Prometheus и Grafana", "Собирают метрики и показывают состояние системы"),
    ], [4.2, 12.5])

    add_heading(doc, "Обработка записи", 1)
    add_numbered(doc, [
        "HAProxy принимает JSON запрос по HTTPS и проверяет API key.",
        "Приложение валидирует поля и проверяет operation_id.",
        "Активный writer подтверждает право записи через lease и fencing epoch.",
        "Пакет фиксируется кворумом NATS JetStream.",
        "Операции попадают в локальный WAL и применяются к состоянию.",
        "История и индексы Redb обновляются в фоне.",
    ])

    add_heading(doc, "Обработка чтения", 1)
    doc.add_paragraph(
        "Чтение выполняется из локальной Redb без сетевого кворума на каждый запрос. "
        "Поддерживаются одиночное и пакетное чтение, поиск по operation_id, выборка по владельцу "
        "или SKU и постраничная навигация. Пока индекс истории восстанавливается, исторические "
        "запросы получают ошибку вместо неполных данных."
    )

    add_heading(doc, "Основные гарантии", 1)
    add_bullets(doc, [
        "Одинаковый operation_id с тем же содержимым считается повтором и не меняет баланс.",
        "Одинаковый operation_id с другим содержимым возвращает HTTP 409.",
        "При потере кворума, writer lease или безопасного доступа к диску запись возвращает HTTP 503.",
        "Один отказавший узел NATS не останавливает запись, поскольку остаётся кворум 2 из 3.",
        "Потерянное локальное состояние приложения можно восстановить из реплицированного журнала.",
    ])

    add_heading(doc, "Измеренные показатели", 1)
    add_table(doc, ["Сценарий", "Результат"], [
        ("Смешанная нагрузка 20 процентов записи и 80 процентов чтения", "495 956 операций в секунду"),
        ("Пакетное чтение", "1,84–1,86 млн операций в секунду"),
        ("Пакетная запись", "около 174 тыс. операций в секунду"),
        ("Переключение приложения через HAProxy", "2 091 мс"),
        ("Потери и дубли при failover лидера журнала", "0"),
    ], [10.5, 6.2])
    p = doc.add_paragraph()
    p.add_run("Ограничение результатов. ").bold = True
    p.add_run(
        "Показатели получены на локальном Docker Desktop с пакетными запросами. "
        "Они характеризуют испытанную конфигурацию, а не гарантируют такую же скорость на другом оборудовании."
    )
    doc.save(OUT / "Краткое описание работы программы.docx")


def build_full():
    doc = base_document(
        "Полное описание виртуального склада",
        "Архитектура, обработка данных, гарантии отказоустойчивости и эксплуатация",
    )
    doc.add_paragraph(
        "Программа реализует отказоустойчивый виртуальный склад. Она принимает операции изменения "
        "остатков, хранит их в реплицированном журнале, строит локальную модель чтения и безопасно "
        "обрабатывает повторные запросы. Основной принцип работы: система подтверждает запись только "
        "после фиксации обязательных данных и возвращает временную ошибку, если безопасное подтверждение невозможно."
    )

    add_heading(doc, "1 Назначение и модель данных", 1)
    doc.add_paragraph(
        "Единицей изменения является операция с уникальным operation_id, владельцем owner_id, кодом товара SKU, "
        "изменением delta и версией события. Текущий остаток рассчитывается для пары owner_id и SKU. "
        "Положительная delta увеличивает остаток, отрицательная уменьшает его."
    )
    add_table(doc, ["Поле", "Назначение"], [
        ("operation_id", "Глобальный идентификатор идемпотентности"),
        ("owner_id", "Владелец складского остатка"),
        ("sku", "Идентификатор товара"),
        ("delta", "Знаковое изменение остатка"),
        ("event_version", "Версия формата бизнес события"),
    ], [4.0, 12.7])
    doc.add_paragraph(
        "Отпечаток операции вычисляется из бизнес полей. Поэтому повтор с тем же идентификатором можно "
        "отличить от ошибочной попытки использовать идентификатор для другого изменения."
    )

    add_heading(doc, "2 Компоненты и связи", 1)
    add_architecture(doc)
    add_table(doc, ["Компонент", "Количество", "Функция"], [
        ("HAProxy 3.2 LTS", "1", "Внешний HTTPS вход, проверка ключа и failover приложений"),
        ("Приложение Rust", "2", "HTTP API, дедупликация, расчёт остатков и read модель"),
        ("NATS JetStream 2.11", "3", "Канонический журнал, кворум и writer lease"),
        ("Redb", "по экземпляру", "Дисковая дедупликация, история и вторичные индексы"),
        ("Prometheus", "1", "Сбор и хранение метрик"),
        ("Grafana", "1", "Панели контроля и наблюдение"),
    ], [4.4, 2.4, 9.9])
    doc.add_paragraph(
        "Все компоненты запускаются через Docker Compose. Два приложения имеют независимые локальные volumes. "
        "Каждый NATS узел также использует отдельный volume, поэтому журнал не зависит от диска приложения."
    )

    add_heading(doc, "3 Путь записи", 1)
    add_numbered(doc, [
        "Клиент отправляет JSON операцию или пакет операций на HAProxy по HTTPS.",
        "HAProxy проверяет API key и направляет запрос приложению, готовому принимать запись.",
        "Приложение проверяет формат, размеры полей, версию события и допустимость delta.",
        "Для operation_id вычисляются хэш и отпечаток содержимого. Проверка выполняется через сегментированный индекс дедупликации.",
        "Writer проверяет действующий lease и получает текущую fencing epoch.",
        "Пакет кодируется во внутренний бинарный формат и публикуется одной записью в NATS JetStream.",
        "JetStream подтверждает запись после выполнения требований репликации. Для кластера из трёх узлов требуется кворум 2 из 3.",
        "Приложение повторно проверяет право writer, записывает операции в локальный WAL и применяет их к соответствующим шардам состояния.",
        "Клиент получает applied или duplicate. Вторичные дисковые индексы и история материализуются асинхронно."
    ])

    add_heading(doc, "4 Параллельная обработка", 1)
    doc.add_paragraph(
        "Состояние распределено между 16 shard workers по хэшу пары owner_id и SKU. Все изменения одной пары "
        "попадают в один worker и сохраняют порядок. Разные пары обрабатываются параллельно. Дедупликационный "
        "индекс разделён на 64 сегмента, чтобы независимые запросы реже ожидали одну общую блокировку."
    )
    doc.add_paragraph(
        "Пакетная обработка объединяет множество операций в один сетевой publish и групповые дисковые записи. "
        "Это повышает пропускную способность, но задержка p95 относится к целому пакету, а не к одной операции."
    )

    add_heading(doc, "5 Идемпотентность и конфликты", 1)
    add_table(doc, ["Ситуация", "Ответ", "Изменение состояния"], [
        ("Новый operation_id", "HTTP 200 applied", "Операция применяется один раз"),
        ("Тот же ID и то же содержимое", "HTTP 200 duplicate", "Баланс не меняется"),
        ("Тот же ID и другое содержимое", "HTTP 409 conflict", "Пакет отклоняется"),
        ("Нет кворума или writer lease", "HTTP 503", "Локальное изменение запрещено"),
    ], [7.2, 4.0, 5.5])
    doc.add_paragraph(
        "Проверка конфликта охватывает весь пакет. Если внутри пакета найдено несовместимое повторное использование "
        "идентификатора, новые операции этого пакета не должны частично попасть в состояние."
    )

    add_heading(doc, "6 Локальное хранение", 1)
    add_heading(doc, "WAL", 2)
    doc.add_paragraph(
        "Write ahead log хранит локально применяемые операции и используется для быстрого восстановления после "
        "аварийного завершения. Записи имеют длину, контрольную сумму CRC32 и бинарную полезную нагрузку. "
        "Незавершённый хвост после аварии отбрасывается, а повреждение завершённой записи обнаруживается."
    )
    add_heading(doc, "Snapshot и checkpoint", 2)
    doc.add_paragraph(
        "Snapshot содержит материальное состояние, номер последовательности и writer epoch. После сохранения snapshot "
        "старый WAL компактируется. Checkpoint позволяет при обычном запуске читать из JetStream только новый хвост. "
        "Snapshot проверяется по SHA 256; повреждённый файл помещается в карантин."
    )
    add_heading(doc, "Redb", 2)
    doc.add_paragraph(
        "Дедупликация и история разделены по независимым файлам Redb. История содержит саму операцию и индексы по "
        "owner_id и SKU. Если локальная история отсутствует, приложение выполняет фоновый backfill из JetStream. "
        "Обычные записи остаются доступными, но историческое чтение не возвращает частичную выборку."
    )

    add_heading(doc, "7 Путь чтения", 1)
    doc.add_paragraph(
        "Чтение не требует quorum commit. Запрос обслуживается локальной read моделью, что отделяет скорость чтения "
        "от сетевой задержки NATS. Доступны текущий баланс, поиск одной операции, пакетный поиск операций, выборки "
        "по владельцу и товару, а также cursor pagination."
    )
    add_bullets(doc, [
        "Пакетный read endpoint открывает одну транзакцию Redb для группы идентификаторов.",
        "Read модель публикует новое поколение целиком, поэтому читатель не видит частично применённый пакет.",
        "Primary и replica ведут собственные read модели из одного канонического журнала.",
        "Метрика history lag показывает отставание индекса от хвоста JetStream."
    ])

    add_heading(doc, "8 Writer lease и защита от split brain", 1)
    doc.add_paragraph(
        "Только одно приложение может подтверждать новые записи. Владение хранится в реплицированном JetStream KV "
        "и обновляется через compare and swap. Каждое новое владение получает монотонную epoch. Маркер epoch также "
        "попадает в канонический журнал, checkpoint и snapshot."
    )
    doc.add_paragraph(
        "Если старый writer возвращается после сетевого разделения, его epoch уже устарела. Такая запись не применяется "
        "ни в live tail, ни при повторном воспроизведении. Событие сохраняется для диагностики в DLQ."
    )

    add_heading(doc, "9 Поведение при отказах", 1)
    add_table(doc, ["Отказ", "Поведение системы"], [
        ("Падение приложения", "HAProxy переводит трафик на второе приложение; новый writer получает новую epoch"),
        ("Падение одного NATS узла", "Кворум 2 из 3 сохраняется; после выборов запись продолжается"),
        ("Падение двух NATS узлов", "Новых записей нет; API возвращает 503 без локального применения"),
        ("Потеря application volume", "Состояние, дедупликация и история перестраиваются из JetStream"),
        ("Повреждение WAL или snapshot", "Файл помещается в карантин, затем выполняется безопасное восстановление"),
        ("Полный локальный диск", "Запись отклоняется; после восстановления повтор распознаётся по operation_id"),
        ("Повреждённое событие", "Событие изолируется в DLQ и не блокирует последующий поток"),
        ("Потеря всех volumes", "Кластер восстанавливается из внешних архивов с проверкой SHA 256"),
    ], [5.3, 11.4])

    add_heading(doc, "10 Retention и сохранность истории", 1)
    doc.add_paragraph(
        "Основной stream использует политику DiscardNew. При достижении лимита новая запись завершается ошибкой, "
        "а старые подтверждённые сообщения не удаляются автоматически ради освобождения места. Это fail closed режим: "
        "оператор должен увеличить ёмкость или выполнить контролируемую процедуру архивирования."
    )
    doc.add_paragraph(
        "Локальная история является восстанавливаемой копией. Каноническим источником остаётся реплицированный stream, "
        "а защита от полной потери кластера обеспечивается внешней резервной копией NATS volumes."
    )

    add_heading(doc, "11 Безопасность", 1)
    add_bullets(doc, [
        "Внешний маршрут использует TLS 1.3.",
        "HAProxy требует API key для прикладных и диагностических маршрутов.",
        "Соединения приложений с NATS используют TLS и отдельные учётные данные.",
        "Прямые диагностические порты привязаны к 127.0.0.1.",
        "В репозитории не хранятся сгенерированные приватные TLS ключи."
    ])
    doc.add_paragraph(
        "Текущие ключи и сертификаты предназначены для локальной демонстрации. В промышленной среде потребуются "
        "внешнее хранилище секретов, регулярная ротация, разделение ролей и сертификаты от доверенного центра."
    )

    add_heading(doc, "12 Наблюдаемость", 1)
    doc.add_paragraph(
        "Prometheus опрашивает оба приложения. Grafana показывает состояние кворума, writer lease, fencing epoch, "
        "последовательности журналов, replication lag, history lag, размеры WAL и Redb, память процесса, ответы 503, "
        "дубли и конфликты. Настроены правила предупреждений по потере кворума, отсутствию writer, split brain, "
        "отставанию репликации и росту памяти."
    )

    add_heading(doc, "13 Производительность", 1)
    add_table(doc, ["Измерение", "Результат"], [
        ("Смешанная нагрузка 20 процентов записи и 80 процентов чтения", "495 956 операций в секунду"),
        ("Первый проход пакетного чтения", "1 857 601 операция в секунду"),
        ("Повторный проход пакетного чтения", "1 838 048 операций в секунду"),
        ("Пакетная запись", "173 999,6 операции в секунду"),
        ("Смешанный p95", "523,574 мс на пакетный запрос"),
        ("Single node baseline", "608 598,1 операции в секунду"),
        ("Quorum failover под нагрузкой", "416 692,1 операции в секунду"),
        ("Переключение приложения", "2 091 мс"),
    ], [10.6, 6.1])
    doc.add_paragraph(
        "Результаты получены на локальном Docker Desktop. Генератор использовал заранее подготовленный набор из "
        "1 000 000 операций, пакетные запросы и конкурентные соединения. Поэтому эти числа следует использовать "
        "для сравнения версий стенда на одной машине, а не как обещание производительности другого окружения."
    )

    add_heading(doc, "14 Основные тесты", 1)
    add_table(doc, ["Файл", "Что проверяет"], [
        ("test-prepared-read-write-performance.ps1", "Скорость записи, чтения и смешанного профиля 20 на 80"),
        ("test-quorum-failover.ps1", "Работу под нагрузкой при остановке лидера JetStream"),
        ("test-application-failover.ps1", "Переключение HAProxy между двумя приложениями"),
        ("test-local-volume-loss.ps1", "Полную перестройку приложения после удаления локального volume"),
    ], [7.8, 8.9])

    add_heading(doc, "15 Запуск", 1)
    doc.add_paragraph("Перед первым запуском создаются локальные сертификаты:")
    p = doc.add_paragraph()
    r = p.add_run(r".\scripts\new-dev-tls-certificate.ps1")
    r.font.name = "Cascadia Mono"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cascadia Mono")
    doc.add_paragraph("Затем собираются и запускаются контейнеры:")
    p = doc.add_paragraph()
    r = p.add_run("docker compose up --build -d")
    r.font.name = "Cascadia Mono"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cascadia Mono")
    add_table(doc, ["Сервис", "Локальный адрес"], [
        ("Публичный HTTPS API", "https://localhost:8443"),
        ("HAProxy HTTP для диагностики", "http://localhost:8080"),
        ("Primary приложение", "http://localhost:8082"),
        ("Replica приложение", "http://localhost:8081"),
        ("Prometheus", "http://localhost:9090"),
        ("Grafana", "http://localhost:3000"),
    ], [7.4, 9.3])

    add_heading(doc, "16 Ограничения", 1)
    add_bullets(doc, [
        "Docker Compose размещает компоненты на одной физической машине и не доказывает устойчивость к потере дата центра.",
        "Нет Kubernetes, автоматического распределения по failure domains и multi region replication.",
        "Redb хранит историю локально и масштабируется хуже отдельного аналитического кластера ClickHouse.",
        "Запись выполняет один активный writer; добавление приложений увеличивает возможности чтения, но не линейно ускоряет запись.",
        "Не проведены длительные soak тесты, независимый аудит безопасности и полноценный fuzzing бинарных форматов.",
        "Демонстрационная конфигурация секретов должна быть заменена перед промышленным использованием."
    ])
    doc.add_paragraph(
        "Таким образом, программа подходит для демонстрации принципов кворумной записи, идемпотентности, failover и "
        "быстрой локальной read модели. Для промышленного внедрения потребуется инфраструктурное разделение узлов, "
        "централизованное управление секретами и длительная проверка на целевом оборудовании."
    )
    doc.save(OUT / "Полное описание работы программы.docx")


if __name__ == "__main__":
    build_short()
    build_full()
    print("created=2")
